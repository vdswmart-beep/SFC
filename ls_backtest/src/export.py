"""Writes the results: one CSV per table, an Excel report and the PNG charts.
The only module with side effects."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from matplotlib.figure import Figure
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

PERCENT_COLUMNS = {"Active return (ann.)", "Tracking error", "Max DD (pocket)", "Alpha (ann.)",
                   "Weight"}
PERCENT_SHEETS = {"Stress tests", "Monthly returns"}  # every number there is a return
PERCENT, RATIO = "0.0%;(0.0%);-", "0.00;(0.00);-"  # finance style: negatives in brackets


def _number_format(value: object, column: str, all_percent: bool) -> str | None:
    """Excel number format of one data cell; None leaves text untouched."""
    if isinstance(value, datetime):
        return "yyyy-mm-dd"
    if not isinstance(value, int | float):
        return None
    if all_percent or column in PERCENT_COLUMNS:
        return PERCENT
    return "0" if column == "Months" else RATIO


def _format_sheet(sheet: object, table: pd.DataFrame, all_percent: bool) -> None:
    """Arial font, bold header, number format by column and readable widths on one sheet."""
    n_index = table.index.nlevels
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.row == 1)
            column = table.columns[cell.column - n_index - 1] if cell.column > n_index else ""
            number_format = _number_format(cell.value, column, all_percent)
            if cell.row > 1 and number_format:
                cell.number_format = number_format
    for k, cells in enumerate(sheet.columns, start=1):
        width = max(len(str(cell.value)) for cell in cells if cell.value is not None)
        sheet.column_dimensions[get_column_letter(k)].width = min(max(width, 8) + 2, 32)
    sheet.freeze_panes = sheet.cell(row=2, column=n_index + 1)


def save_outputs(tables: dict[str, pd.DataFrame], figures: dict[str, Figure],
                 output_dir: str) -> None:
    """Writes each table as CSV, all tables in report.xlsx (one sheet each) and each figure as
    PNG, in output_dir (created if needed)."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    for name, table in tables.items():
        table.to_csv(out / f"{name.lower().replace(' ', '_')}.csv")
    with pd.ExcelWriter(out / "report.xlsx", engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, merge_cells=False)
            _format_sheet(writer.sheets[name], table, name in PERCENT_SHEETS)
    for name, figure in figures.items():
        figure.savefig(out / name, dpi=200, bbox_inches="tight")
